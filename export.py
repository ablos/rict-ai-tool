"""Write claims as JSONL (for machines) and as Excel review workbooks (for people).

The review is blind: reviewers first see only what the pipeline found (quotes with
their context), decide themselves what each quote says, and only then compare with
what the LLM decided. A review workbook covers one paper and has these sheets:

    Judge      a random sample of quotes, with one column per relation; the reviewer puts
               an x under every relation the quote supports and writes the value
    Relations  the ontology's relations and definitions, for reference while judging
    Compare    the reviewer's relations next to the LLM's (open only after judging);
               the reviewer marks whether the LLM's values are right
    Summary    agreement between reviewer and LLM, overall and per relation
    Run info   the settings of the run
"""

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VALUE_OK = ["yes", "partly", "no"]

# The Judge sheet: these columns, then one column per relation, then the review columns.
JUDGE_COLUMNS = ["item_id", "method", "quote", "context", "page", "section"]
SUPPORTS_NOTHING = "supports_nothing"  # x here: the quote supports no relation for this method
JUDGE_REVIEW_COLUMNS = [SUPPORTS_NOTHING, "human_value", "notes", "reviewer"]
# The Compare sheet: these visible columns, then hidden helper columns for the formulas.
# llm_only_correct: of the llm_only relations, how many are right after all (the reviewer missed them).
COMPARE_COLUMNS = ["item_id", "quote", "human_relations", "llm_relations", "agreed", "llm_only",
                   "llm_only_correct", "human_only", "human_value", "llm_values", "value_ok", "notes"]
COMPARE_HELPERS = ["judged", "llm_count", "human_count", "agreed_count", "exact", "exact_corrected"]

# Column widths in characters; columns not listed get DEFAULT_WIDTH.
WIDTHS = {
    "method": 20, "quote": 60, "context": 80, "section": 26, "human_value": 30, "notes": 30,
    "human_relations": 26, "llm_relations": 26, "agreed": 26, "llm_only": 22, "llm_only_correct": 18, "human_only": 22,
    "llm_values": 40, "value_ok": 10,
    "relation": 22, "concept": 22, "definition": 90,
    "paper": 40, "human_decision": 16, "human_methods": 30, "llm_decision": 14, "decision_match": 16, "criteria_not_met": 18,
    "llm_included": 30, "llm_rejected": 40,
}
DEFAULT_WIDTH = 12

GREY = PatternFill("solid", fgColor="E7E6E6")  # filled in automatically, not by the reviewer


def write_jsonl(rows: list[dict], path: str | Path) -> None:
    with open(path, "w", encoding="utf-8") as file:
        for row in rows:
            file.write(json.dumps(row, ensure_ascii=False) + "\n")


def read_jsonl(path: str | Path) -> list[dict]:
    with open(path, encoding="utf-8") as file:
        return [json.loads(line) for line in file if line.strip()]


def review_items(rows: list[dict]) -> list[dict]:
    """The claims grouped by quote: one item per (method, quote), with every LLM claim
    that quote supports. Reviewers judge quotes, so a quote the LLM used for three
    values is one item, not three."""
    items = {}

    for row in rows:
        if row["status"] != "reported":
            continue

        key = (row["method"], row["quote"])
        item = items.setdefault(key, {
            "method": row["method"], "quote": row["quote"], "context": row.get("context"),
            "page": row.get("page"), "section": row.get("section"), "claims": [],
        })
        item["claims"].append({"claim_id": row["claim_id"], "relation": row["relation"], "value": row["value"]})

    return list(items.values())


def sample_items(items: list[dict], size: int, seed: int) -> list[dict]:
    """A random sample spread over methods and relations: take one item from each
    (method, relation) group in turn, so every method and rare relations are checked too.
    size 0 means all items."""
    if size <= 0 or size >= len(items):
        return items

    rng = random.Random(seed)
    groups = {}

    for item in items:
        groups.setdefault((item["method"], item["claims"][0]["relation"]), []).append(item)

    groups = list(groups.values())

    for group in groups:
        rng.shuffle(group)

    sample = []

    while len(sample) < size:
        # New random order each pass, so with more groups than places the choice is random too.
        rng.shuffle(groups)

        for group in groups:
            if group and len(sample) < size:
                sample.append(group.pop())

    rng.shuffle(sample)  # so reviewers don't see the relations in a pattern
    return sample


def write_review_sheet(rows: list[dict], path: str | Path, ontology: dict,
                       sample_size: int, seed: int, meta: dict | None = None) -> None:
    """A blind review workbook for the given rows (normally the rows of one paper)."""
    items = sample_items(review_items(rows), sample_size, seed)

    for number, item in enumerate(items, start=1):
        item["item_id"] = f"Q{number:02d}"

    relation_names = [relation["name"] for concept in ontology["concepts"] for relation in concept["relations"]]
    labels = {relation["name"]: relation_label(concept, relation)
              for concept in ontology["concepts"] for relation in concept["relations"]}
    workbook = Workbook()

    judge = workbook.active
    judge.title = "Judge"
    judge_headers = JUDGE_COLUMNS + [labels[name] for name in relation_names] + JUDGE_REVIEW_COLUMNS
    _fill_sheet(judge, [[item.get(column) for column in JUDGE_COLUMNS] for item in items], judge_headers)

    # The relation columns: the label at an angle so it stays readable, and an "x" dropdown.
    # The x columns run from the first relation up to and including supports_nothing.
    first, last = _letter(judge_headers, labels[relation_names[0]]), _letter(judge_headers, SUPPORTS_NOTHING)
    judge.column_dimensions[last].width = 7
    judge[f"{last}1"].alignment = Alignment(text_rotation=45, vertical="bottom")

    for name in relation_names:
        judge.column_dimensions[_letter(judge_headers, labels[name])].width = 7
        judge[f"{_letter(judge_headers, labels[name])}1"].alignment = Alignment(text_rotation=45, vertical="bottom")

    # Room for the longest angled label, e.g. "addressesValue (Ethical or Sustainability Value)".
    # At 45 degrees a label needs about 0.7 x its length in height: ~4 points per character.
    judge.row_dimensions[1].height = 4 * max(len(label) for label in labels.values()) + 10

    for cells in judge[f"{first}2:{last}{max(len(items) + 1, 2)}"]:
        for cell in cells:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    dropdown = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dropdown.add(f"{first}2:{last}{max(len(items) + 1, 2)}")
    judge.add_data_validation(dropdown)

    # Narrower text columns, and keep id, method and quote visible while scrolling to the relations.
    for name, width in (("quote", 45), ("context", 55), ("section", 18)):
        judge.column_dimensions[_letter(judge_headers, name)].width = width

    judge.freeze_panes = f"{_letter(judge_headers, 'context')}2"

    reference = workbook.create_sheet("Relations")
    _fill_sheet(reference, [[relation["name"], concept["label"], relation["definition"]]
                            for concept in ontology["concepts"] for relation in concept["relations"]],
                ["relation", "concept", "definition"])

    compare = workbook.create_sheet("Compare")
    _add_compare(compare, items, relation_names, judge_headers, labels)

    summary = workbook.create_sheet("Summary")
    _add_summary(summary, relation_names, len(items), labels)

    if meta:
        info = workbook.create_sheet("Run info")

        for key, value in _flatten(meta):
            info.append([key, value])

        info.column_dimensions["A"].width = 30
        info.column_dimensions["B"].width = 80

    workbook.save(path)


DECISIONS = ["include", "exclude"]
NO_METHOD_FOUND = "(no method found)"  # the row for a paper where the LLM found no candidate method
SCREEN_JUDGE_COLUMNS = ["paper", "method", "human_decision", "notes", "reviewer"]
SCREEN_COMPARE_COLUMNS = ["paper", "method", "human_decision", "llm_decision", "decision_match",
                          "criteria_not_met", "llm_reasons", "notes"]


def write_screening_sheet(papers: list[dict], path: str | Path, meta: dict | None = None) -> None:
    """A blind check of the screening step: one row per candidate method the LLM found.

    Judge     the reviewer decides include/exclude per method (does it meet all criteria?)
              without seeing the LLM's decision; methods the LLM missed can be added as rows
    Compare   the LLM's decision, its reason per criterion, and whether it matches
    Summary   how often reviewer and LLM agree, and which way they disagree
    """
    rows = []

    for paper in papers:
        name = Path(paper["pdf"]).name

        for candidate in paper["screening"]:
            failed = [criterion for criterion, check in candidate["criteria"].items() if not check["met"]]
            reasons = "\n".join(f"{criterion} {'met' if check['met'] else 'NOT met'}: {check['reason']}"
                                for criterion, check in candidate["criteria"].items())
            rows.append([name, candidate["name"], "include" if candidate["included"] else "exclude",
                         ", ".join(failed), reasons])

        if not paper["screening"]:
            rows.append([name, NO_METHOD_FOUND, "exclude", "", "The LLM found no candidate method in this paper."])

    workbook = Workbook()
    judge = workbook.active
    judge.title = "Judge"
    _fill_sheet(judge, [[paper, method] for paper, method, *_ in rows], SCREEN_JUDGE_COLUMNS)
    # Room below the rows, so reviewers can add methods the LLM missed.
    _add_dropdown(judge, SCREEN_JUDGE_COLUMNS, "human_decision", f'"{",".join(DECISIONS)}"', len(rows) + 50)

    compare = workbook.create_sheet("Compare")
    human = _letter(SCREEN_JUDGE_COLUMNS, "human_decision")
    data = [[paper, method, f'=IF(Judge!{human}{row_number}="","",Judge!{human}{row_number})',
             llm_decision, None, failed, reasons, None]
            for row_number, (paper, method, llm_decision, failed, reasons) in enumerate(rows, start=2)]
    _fill_sheet(compare, data, SCREEN_COMPARE_COLUMNS)
    human_c, llm_c, match_c = (_letter(SCREEN_COMPARE_COLUMNS, name) for name in ("human_decision", "llm_decision", "decision_match"))

    for row_number in range(2, len(rows) + 2):
        compare[f"{match_c}{row_number}"] = (f'=IF({human_c}{row_number}="","",'
                                             f'IF({human_c}{row_number}={llm_c}{row_number},"yes","no"))')

        for column in (human_c, match_c):
            compare[f"{column}{row_number}"].fill = GREY

    compare.column_dimensions[_letter(SCREEN_COMPARE_COLUMNS, "llm_reasons")].width = 100

    summary = workbook.create_sheet("Summary")
    last = len(rows) + 1
    human_r, llm_r, match_r = (f"Compare!${c}$2:${c}${last}" for c in (human_c, llm_c, match_c))
    judge_method = f"Judge!$B${last + 1}:$B$1000"
    judge_decision = f"Judge!${human}${last + 1}:${human}$1000"
    summary.append(["Measure (per method)", "Value"])
    summary.append(["Methods found by the LLM", len(rows)])
    summary.append(["Methods judged", f'=COUNTIF({match_r},"yes")+COUNTIF({match_r},"no")'])
    summary.append(["Same decision as the LLM", f'=IFERROR(COUNTIF({match_r},"yes")/B3,"")'])
    summary.append(["LLM included, reviewer excluded", f'=COUNTIFS({llm_r},"include",{human_r},"exclude")'])
    summary.append(["LLM excluded, reviewer included", f'=COUNTIFS({llm_r},"exclude",{human_r},"include")'])
    summary.append(["Methods added by the reviewer (missed by the LLM)",
                    f'=COUNTIFS({judge_method},"<>",{judge_decision},"include")'])
    summary["B4"].number_format = "0%"
    summary.column_dimensions["A"].width = 50

    for cell in summary[1]:
        cell.font = Font(bold=True)

    if meta:
        info = workbook.create_sheet("Run info")

        for key, value in _flatten(meta):
            info.append([key, value])

        info.column_dimensions["A"].width = 30
        info.column_dimensions["B"].width = 80

    workbook.save(path)


def _compare_headers(relation_names: list[str]) -> list[str]:
    """Visible columns, then hidden helpers: 0/1 per relation for the reviewer (H) and the LLM (L)."""
    return (COMPARE_COLUMNS + COMPARE_HELPERS
            + [f"H_{name}" for name in relation_names] + [f"L_{name}" for name in relation_names])


def _add_compare(sheet, items: list[dict], relation_names: list[str], judge_headers: list[str],
                 labels: dict[str, str]) -> None:
    """Per quote: the reviewer's relations (read from Judge by formula) next to the LLM's.

    Hidden helper columns hold a 0/1 per relation for both sides, so agreement can be
    counted with plain SUM and SUMPRODUCT formulas.
    """
    headers = _compare_headers(relation_names)
    col = {name: _letter(headers, name) for name in headers}
    reviewer = _letter(judge_headers, "reviewer")
    human_value = _letter(judge_headers, "human_value")
    marks_first = _letter(judge_headers, labels[relation_names[0]])
    marks_last = _letter(judge_headers, SUPPORTS_NOTHING)
    data = []

    for r, item in enumerate(items, start=2):
        llm = {claim["relation"] for claim in item["claims"]}
        h = {name: f"{col['H_' + name]}{r}" for name in relation_names}
        l = {name: f"{col['L_' + name]}{r}" for name in relation_names}
        h_range = f"{col['H_' + relation_names[0]]}{r}:{col['H_' + relation_names[-1]]}{r}"
        l_range = f"{col['L_' + relation_names[0]]}{r}:{col['L_' + relation_names[-1]]}{r}"
        judged = f"{col['judged']}{r}"

        def names_where(condition) -> str:
            """A formula listing the relation names for which condition(name) holds."""
            parts = "&".join(f'IF({condition(name)},"{name}; ","")' for name in relation_names)
            return f'=IF({judged}=0,"",{parts})'

        values = {
            "item_id": item["item_id"],
            "quote": item["quote"],
            "human_relations": names_where(lambda name: f"{h[name]}=1"),
            "llm_relations": "; ".join(name for name in relation_names if name in llm),
            "agreed": names_where(lambda name: f"AND({h[name]}=1,{l[name]}=1)"),
            "llm_only": names_where(lambda name: f"AND({h[name]}=0,{l[name]}=1)"),
            "human_only": names_where(lambda name: f"AND({h[name]}=1,{l[name]}=0)"),
            "human_value": f'=IF(Judge!{human_value}{r}="","",Judge!{human_value}{r})',
            "llm_values": "\n".join(f"{claim['relation']}: {claim['value']}" for claim in item["claims"]),
            # Judged as soon as the row has anything: an x, a value or a reviewer name.
            "judged": f'=IF(OR(COUNTA(Judge!{marks_first}{r}:{marks_last}{r})>0,TRIM(Judge!{human_value}{r})<>"",'
                      f'TRIM(Judge!{reviewer}{r})<>""),1,0)',
            "llm_count": f"=SUM({l_range})",
            "human_count": f"=SUM({h_range})",
            "agreed_count": f"=SUMPRODUCT({h_range},{l_range})",
            "exact": f'=IF(AND({judged}=1,{col["llm_count"]}{r}={col["human_count"]}{r},'
                     f'{col["agreed_count"]}{r}={col["llm_count"]}{r}),1,0)',
            # Corrected: every reviewer relation was found, and all LLM-only relations are right.
            "exact_corrected": f'=IF(AND({judged}=1,{col["human_count"]}{r}={col["agreed_count"]}{r},'
                               f'{col["llm_count"]}{r}-{col["agreed_count"]}{r}=N({col["llm_only_correct"]}{r})),1,0)',
        }

        for name in relation_names:
            values[f"H_{name}"] = f'=IF(TRIM(Judge!{_letter(judge_headers, labels[name])}{r})="",0,1)'
            values[f"L_{name}"] = 1 if name in llm else 0

        data.append([values.get(name) for name in headers])

    _fill_sheet(sheet, data, headers)

    for r in range(2, len(items) + 2):
        for name in ("human_relations", "agreed", "llm_only", "human_only", "human_value"):
            sheet[f"{col[name]}{r}"].fill = GREY

    for name in COMPARE_HELPERS + [h for h in headers if h.startswith(("H_", "L_"))]:
        sheet.column_dimensions[col[name]].hidden = True

    _add_dropdown(sheet, headers, "value_ok", f'"{",".join(VALUE_OK)}"', len(items))

    # llm_only_correct takes a whole number from 0 up.
    counts = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    counts.add(f"{col['llm_only_correct']}2:{col['llm_only_correct']}{max(len(items) + 1, 2)}")
    sheet.add_data_validation(counts)


def _add_summary(sheet, relation_names: list[str], item_count: int, labels: dict[str, str]) -> None:
    """Agreement between reviewer and LLM, overall and per relation.

    Relation precision: of the relations the LLM gave the quotes, how many the reviewer also gave.
    Relation recall:    of the relations the reviewer gave, how many the LLM also gave.
    Corrected versions also count the LLM-only relations the reviewer marked as right after all.
    """
    headers = _compare_headers(relation_names)
    last = item_count + 1

    def span(name: str) -> str:
        return f"Compare!${_letter(headers, name)}$2:${_letter(headers, name)}${last}"

    judged, llm_count, human_count, agreed, exact, value_ok, correct, exact_corrected = (
        span(name) for name in ("judged", "llm_count", "human_count", "agreed_count", "exact", "value_ok",
                                "llm_only_correct", "exact_corrected"))

    sheet.append(["Measure", "Value", "± (95%)"])
    sheet.append(["Quotes in the sample", item_count])
    sheet.append(["Quotes judged", f"=SUM({judged})"])
    sheet.append(["Exactly the same relations as the LLM", f'=IFERROR(SUM({exact})/B3,"")',
                  '=IFERROR(1.96*SQRT(B4*(1-B4)/B3),"")'])
    sheet.append(["Relation precision (LLM's relations the reviewer agrees with)",
                  f'=IFERROR(SUMPRODUCT({judged},{agreed})/SUMPRODUCT({judged},{llm_count}),"")'])
    sheet.append(["Relation recall (reviewer's relations the LLM also found)",
                  f'=IFERROR(SUMPRODUCT({judged},{agreed})/SUMPRODUCT({judged},{human_count}),"")'])
    sheet.append(["Values right (quotes with at least one agreed relation)",
                  f'=IFERROR(COUNTIF({value_ok},"yes")/COUNTIFS({agreed},">0",{judged},1),"")'])
    sheet.append(["Reviewer: quote supports nothing about the method",
                  f'=IFERROR(COUNTIFS({judged},1,{human_count},0)/B3,"")'])
    # Rows 9-13: corrected with llm_only_correct, for relations the LLM found and the reviewer missed.
    sheet.append(["Corrected for relations the reviewer missed (llm_only_correct):"])
    sheet.append(["LLM-only relations that were right after all",
                  f'=IFERROR(SUMPRODUCT({judged},{correct})/(SUMPRODUCT({judged},{llm_count})-SUMPRODUCT({judged},{agreed})),"")'])
    sheet.append(["Exactly the same relations (corrected)", f'=IFERROR(SUM({exact_corrected})/B3,"")',
                  '=IFERROR(1.96*SQRT(B11*(1-B11)/B3),"")'])
    sheet.append(["Relation precision (corrected)",
                  f'=IFERROR((SUMPRODUCT({judged},{agreed})+SUMPRODUCT({judged},{correct}))/SUMPRODUCT({judged},{llm_count}),"")'])
    sheet.append(["Relation recall (corrected)",
                  f'=IFERROR((SUMPRODUCT({judged},{agreed})+SUMPRODUCT({judged},{correct}))'
                  f'/(SUMPRODUCT({judged},{human_count})+SUMPRODUCT({judged},{correct})),"")'])
    sheet.append([])
    sheet.append(["Relation (blind, not corrected)", "LLM chose", "Reviewer chose", "Both", "Precision", "Recall"])
    header_row = sheet.max_row

    for relation in relation_names:
        row = sheet.max_row + 1
        sheet.append([
            labels[relation],
            f"=SUMPRODUCT({judged},{span('L_' + relation)})",
            f"=SUMPRODUCT({judged},{span('H_' + relation)})",
            f"=SUMPRODUCT({judged},{span('H_' + relation)},{span('L_' + relation)})",
            f'=IFERROR(D{row}/B{row},"")',
            f'=IFERROR(D{row}/C{row},"")',
        ])

    for row_number in (1, header_row):
        for cell in sheet[row_number]:
            cell.font = Font(bold=True)

    sheet.cell(9, 1).font = Font(bold=True)

    for row_number in [*range(4, 9), *range(10, 14)]:
        for column in (2, 3):
            sheet.cell(row_number, column).number_format = "0%"

    for row_number in range(header_row + 1, sheet.max_row + 1):
        for column in (5, 6):
            sheet.cell(row_number, column).number_format = "0%"

    sheet.column_dimensions["A"].width = 62

    for column in "BCDEF":
        sheet.column_dimensions[column].width = 15


def relation_label(concept: dict, relation: dict) -> str:
    """The Judge column header for a relation, e.g. "addresses (Problem Type)"."""
    return f"{relation['name']} ({concept['label']})"


def relation_from_label(label: str) -> str:
    """The relation name back from a Judge column header."""
    return label.split(" (")[0]


def _letter(headers: list[str], name: str) -> str:
    return get_column_letter(headers.index(name) + 1)


def _fill_sheet(sheet, data: list[list], headers: list[str]) -> None:
    """A header row and the data; columns without data stay empty for the reviewer."""
    sheet.append(headers)

    for values in data:
        sheet.append([_cell_value(value) for value in values] + [None] * (len(headers) - len(values)))

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for index, name in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(name, DEFAULT_WIDTH)

    sheet.freeze_panes = "B2"  # header row and first column stay visible while scrolling
    sheet.auto_filter.ref = sheet.dimensions


def _add_dropdown(sheet, headers: list[str], column: str, source: str, rows: int) -> None:
    """source is a quoted list ('"yes,no"') or a cell range ("Relations!$A$2:$A$9")."""
    letter = _letter(headers, column)
    dropdown = DataValidation(type="list", formula1=source, allow_blank=True)
    dropdown.add(f"{letter}2:{letter}{max(rows + 1, 2)}")
    sheet.add_data_validation(dropdown)


def _cell_value(value):
    """Excel cells hold text, numbers and booleans; anything else is written as JSON."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value

    return json.dumps(value, ensure_ascii=False)


def _flatten(data: dict, prefix: str = ""):
    """Yield ("a.b", value) pairs from a nested dict, for the Run info sheet."""
    for key, value in data.items():
        if isinstance(value, dict):
            yield from _flatten(value, f"{prefix}{key}.")

        else:
            yield f"{prefix}{key}", _cell_value(value)
